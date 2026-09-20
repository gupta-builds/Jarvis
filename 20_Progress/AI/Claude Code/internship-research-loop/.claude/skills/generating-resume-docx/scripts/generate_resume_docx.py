#!/usr/bin/env python3
"""Generate a resume .docx that matches this skill's format rules (cross-referenced
from Resume Alteration Standard §8, not re-researched -- see reference/resume-format-rules.md).

Format only, same division of labor as generate_cover_letter_docx.py: this script
never decides which bullets appear or in what order -- that's Resume Alteration
Standard §2/§3's job and .cursor/skills/resume-alteration's job. This script turns
an already-approved bullet selection into a real, single-column, ATS-safe .docx.

Rules implemented, each cited in reference/resume-format-rules.md:
  - Single column, no tables/text boxes/graphics (Greenhouse's own parse-failure list).
  - One font throughout, 10-12pt, standard system font.
  - Real Word bullet-list style for every bullet, not a typed hyphen.
  - Bold section headings one size larger than body; bold sub-entry headers
    (role/company/dates, project names) at body size -- matches Main Resume.md's
    own real visual hierarchy.
  - One-page soft target (word-count warning, not a hard block -- a two-page
    resume is a legitimate per-application choice for technical roles where the
    target company's own guidance allows it, per the Standard's own §8 note).
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass, field

try:
    from docx import Document
    from docx.shared import Inches, Pt
except ImportError:
    print(
        "python-docx is not installed. Install it (python-docx==1.2.0, matching "
        "internship-research-loop's requirements.txt) before running this for real.",
        file=sys.stderr,
    )
    raise

APPROVED_FONTS = {"Times New Roman", "Arial", "Calibri", "Garamond", "Helvetica"}
MIN_FONT_SIZE = 10
MAX_FONT_SIZE = 12
# One-page soft ceiling. A two-page resume is a legitimate, per-application choice
# for technical roles (see reference/resume-format-rules.md) -- this is a warning
# the human weighs, never a hard block.
ONE_PAGE_WORD_CEILING = 550


@dataclass
class Entry:
    """One Experience/Project sub-entry: a header line plus its real bullets."""

    header: str  # e.g. "Software Engineering Intern — Acme Robotics, June-August 2026"
    bullets: list[str] = field(default_factory=list)


@dataclass
class ResumePlan:
    """Everything this script needs -- no defaults on the required fields, so a
    caller can't accidentally pass through unapproved/placeholder content."""

    name: str
    contact_line: str
    education: str
    skills: list[tuple[str, str]]  # [(category, comma-separated items), ...]
    experience: list[Entry]
    projects: list[Entry] = field(default_factory=list)
    certifications: list[str] = field(default_factory=list)
    objective: str | None = None
    font: str = "Times New Roman"
    font_size: int = 11


def _validate(plan: ResumePlan) -> list[str]:
    warnings: list[str] = []

    if plan.font not in APPROVED_FONTS:
        warnings.append(f"Font '{plan.font}' is not in the Standard-approved set ({', '.join(sorted(APPROVED_FONTS))}).")
    if not (MIN_FONT_SIZE <= plan.font_size <= MAX_FONT_SIZE):
        warnings.append(f"Font size {plan.font_size}pt is outside the sourced 10-12pt range.")

    word_count = len(plan.education.split())
    if plan.objective:
        word_count += len(plan.objective.split())
    for _, items in plan.skills:
        word_count += len(items.split())
    for entry in [*plan.experience, *plan.projects]:
        word_count += len(entry.header.split())
        word_count += sum(len(b.split()) for b in entry.bullets)
    word_count += sum(len(c.split()) for c in plan.certifications)

    if word_count > ONE_PAGE_WORD_CEILING:
        warnings.append(
            f"~{word_count} words -- likely over one page at {plan.font_size}pt. "
            "A second page is a legitimate choice for a technical role where the "
            "target company's own guidance explicitly allows it (Standard §8) -- "
            "this is a flag for that decision, not an error."
        )

    return warnings


def build_resume(plan: ResumePlan, output_path: str) -> list[str]:
    warnings = _validate(plan)

    document = Document()
    section = document.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = plan.font
    normal.font.size = Pt(plan.font_size)

    def heading_line(text: str) -> None:
        p = document.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(plan.font_size + 1)
        p.paragraph_format.space_before = Pt(plan.font_size)
        p.paragraph_format.space_after = Pt(plan.font_size * 0.5)

    def sub_header(text: str) -> None:
        p = document.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        p.paragraph_format.space_before = Pt(plan.font_size * 0.5)

    def bullet(text: str) -> None:
        document.add_paragraph(text, style="List Bullet")

    def body_line(text: str) -> None:
        document.add_paragraph(text)

    # Header: name (bold, larger) then contact line -- same shape as the cover
    # letter's header/closing, so both documents present identical contact info.
    p = document.add_paragraph()
    run = p.add_run(plan.name)
    run.bold = True
    run.font.size = Pt(plan.font_size + 3)
    body_line(plan.contact_line)

    heading_line("Education")
    body_line(plan.education)

    if plan.objective:
        heading_line("Objective")
        body_line(plan.objective)

    heading_line("Skills")
    for category, items in plan.skills:
        p = document.add_paragraph()
        cat_run = p.add_run(f"{category}: ")
        cat_run.italic = True
        p.add_run(items)

    if plan.experience:
        heading_line("Experience")
        for entry in plan.experience:
            sub_header(entry.header)
            for b in entry.bullets:
                bullet(b)

    if plan.projects:
        heading_line("Projects")
        for entry in plan.projects:
            sub_header(entry.header)
            for b in entry.bullets:
                bullet(b)

    if plan.certifications:
        heading_line("Certifications")
        for c in plan.certifications:
            bullet(c)

    document.save(output_path)
    return warnings


def _demo_plan() -> ResumePlan:
    contact = "612-555-0143 · jordan.lee@example.edu · linkedin.com/in/jordanlee-example · github.com/jlee-example"
    return ResumePlan(
        name="Jordan Lee",
        contact_line=contact,
        education="Pursuing a Bachelor of Science in Computer Science — Example State University. Expected Spring 2028.",
        skills=[
            ("Programming", "Python, Rust, TypeScript, JavaScript"),
            ("AI & Data", "LLM APIs, RAG, embeddings, data pipelines"),
            ("Full Stack", "Next.js, React, REST APIs, backend logic"),
        ],
        experience=[
            Entry(
                header="Software Engineering Intern — Acme Robotics, June-August 2026",
                bullets=[
                    "Built a real-time perception dashboard using Rust and a WebSocket event stream.",
                    "Reduced pipeline stall detection time from minutes to seconds by adding structured observability tooling.",
                ],
            ),
            Entry(
                header="Web Development Intern — Example Corp, June-August 2025",
                bullets=["Built and deployed a production web application with Next.js, React, and a Strapi backend."],
            ),
        ],
        projects=[
            Entry(
                header="Second Brain Note Visualizer",
                bullets=["Added a computer-vision layer to an Obsidian vault to visualize idea connections."],
            ),
        ],
        certifications=["Generative AI — Example Learning Platform, issued 2026."],
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--demo", action="store_true", help="Build the fabricated example resume as a self-check.")
    parser.add_argument("--out", default="demo-resume.docx", help="Output .docx path.")
    args = parser.parse_args()

    if not args.demo:
        parser.error("This CLI entry point only runs the --demo self-check. Call build_resume() "
                      "directly from an agent/skill with a real, human-approved ResumePlan for real use.")

    warnings = build_resume(_demo_plan(), args.out)
    print(f"Wrote {args.out}")
    if warnings:
        print("Format warnings:")
        for w in warnings:
            print(f"  - {w}")
    else:
        print("No format warnings -- matches every sourced rule in reference/resume-format-rules.md.")
